import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

prs = Presentation(r"ECDAT_SIH2026_PPT (winner1).pptx")
for i, slide in enumerate(prs.slides):
    print(f"=== SLIDE {i+1} ===")
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            print(shape.text.strip()[:600])
    print()
